"""
BSM 401 Bitirme Kitapçığı üreteci.

MNACP projesi için Sakarya Üniversitesi şablonuna uygun DOCX üretir.
Times New Roman, 12 punto gövde, 14 punto kalın başlıklar.
"""
from __future__ import annotations

import os
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUTPUT = Path(__file__).parent / "MNACP_Bitirme_Kitapcik_v3.docx"
PROJECT_ROOT = Path(__file__).parent.parent
EVAL_DIR = PROJECT_ROOT / "mnacp" / "evaluation" / "results"

PLACEHOLDER_OGRENCI_NO = "[Öğrenci No]"
PLACEHOLDER_ADSOYAD = "[Ad SOYAD]"
PLACEHOLDER_DANISMAN = "[Prof./Doç./Dr. Ad SOYAD]"
PLACEHOLDER_DONEM = "[2025-2026 Bahar Dönemi]"


# ---------- Yardımcılar ----------

def add_page_number_field(paragraph) -> None:
    """Paragrafa { PAGE } alanı ekler."""
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = "PAGE"
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)


def set_default_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:cs"), "Times New Roman")
    pf = style.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)


def set_margins(section, top=2.5, right=2.5, bottom=2.5, left=3.5) -> None:
    section.top_margin = Cm(top)
    section.right_margin = Cm(right)
    section.bottom_margin = Cm(bottom)
    section.left_margin = Cm(left)


def add_centered_paragraph(doc, text, *, bold=False, size=12, space_before=0, space_after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = "Times New Roman"
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_heading_1(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    run.font.name = "Times New Roman"
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.keep_with_next = True
    p.style = doc.styles["Heading 1"] if "Heading 1" in [s.name for s in doc.styles] else p.style
    return p


def add_subheading(doc, text, level=2):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    run.font.name = "Times New Roman"
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    return p


def add_body(doc, text, justify=True, indent=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.size = Pt(12)
    run.font.name = "Times New Roman"
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(12)
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.25)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    run = p.runs[0] if p.runs else p.add_run()
    run.text = ""
    p.add_run(text).font.size = Pt(12)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    return p


def add_caption_under(doc, text, *, size=9):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(12)
    return p


def add_caption_over(doc, text, *, size=9):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(2)
    return p


def add_image_centered(doc, path: Path, *, width_cm=14):
    if not path.exists():
        add_body(doc, f"[Görsel bulunamadı: {path.name}]")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))


def add_table_simple(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(11)
        run.font.name = "Times New Roman"
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for r, row_data in enumerate(rows, start=1):
        for c, val in enumerate(row_data):
            cell = table.rows[r].cells[c]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            run.font.size = Pt(11)
            run.font.name = "Times New Roman"
    return table


def add_page_break(doc):
    doc.add_page_break()


# ---------- KAPAK ----------

def build_cover(doc):
    # Üstten boşluk
    for _ in range(2):
        doc.add_paragraph()

    add_centered_paragraph(doc, "T.C.", bold=True, size=14)
    add_centered_paragraph(doc, "SAKARYA ÜNİVERSİTESİ", bold=True, size=14)
    add_centered_paragraph(
        doc, "BİLGİSAYAR VE BİLİŞİM BİLİMLERİ FAKÜLTESİ", bold=True, size=14, space_after=24,
    )
    add_centered_paragraph(doc, "BSM 401 BİLGİSAYAR MÜHENDİSLİĞİ TASARIMI", bold=True, size=14, space_after=48)

    add_centered_paragraph(
        doc,
        "LLM AJANLARI İÇİN DİNAMİK KEŞİF VE",
        bold=True, size=14,
    )
    add_centered_paragraph(
        doc,
        "DELEGASYON PROTOKOLÜ (MNACP)",
        bold=True, size=14, space_after=48,
    )

    add_centered_paragraph(doc, f"{PLACEHOLDER_OGRENCI_NO} - {PLACEHOLDER_ADSOYAD}", bold=True)
    doc.add_paragraph()
    doc.add_paragraph()

    # Bölüm / Danışman
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Bölüm\t\t:\tBİLGİSAYAR MÜHENDİSLİĞİ").bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Danışman\t:\t{PLACEHOLDER_DANISMAN}").bold = True

    for _ in range(6):
        doc.add_paragraph()
    add_centered_paragraph(doc, PLACEHOLDER_DONEM, bold=True)
    add_page_break(doc)


# ---------- ÖN SAYFALAR ----------

def build_onsoz(doc):
    add_heading_1(doc, "ÖNSÖZ")
    add_body(
        doc,
        "Bu çalışma, Sakarya Üniversitesi Bilgisayar Mühendisliği Bölümü BSM 401 "
        "Bilgisayar Mühendisliği Tasarımı dersi kapsamında hazırlanmış bitirme "
        "projesidir. Çalışmanın motivasyonu büyük dil modellerinin (LLM) tek "
        "başlarına çözemediği çok adımlı görevleri, birbirinin yeteneklerini "
        "çalışma zamanında keşfeden ve birbirine görev devredebilen birden fazla "
        "ajanın koordinasyonu ile çözmektir. Anthropic'in Kasım 2024'te yayınladığı "
        "Model Context Protocol (MCP) standardı tek bir LLM'in araç sunucularına "
        "bağlanmasını tanımlar; ancak ajanların birbirini araç olarak kullanması "
        "kapsamı dışındadır. Bu projede MCP'nin bu boşluğunu dolduran bir "
        "koordinasyon protokolü ve referans uygulaması geliştirilmiştir.",
    )
    add_body(
        doc,
        "Tasarım sürecinde bana yol gösteren danışman hocama, geliştirme "
        "boyunca Claude API'sini sağlayan Anthropic ekosistemine ve açık "
        "kaynak yazılım topluluğuna teşekkürlerimi sunarım.",
    )
    add_page_break(doc)


def build_icindekiler_placeholder(doc):
    add_heading_1(doc, "İÇİNDEKİLER")
    add_body(
        doc,
        "(Bu sayfa Word tarafında otomatik içindekiler oluşturulacaktır. "
        "Şablonun gerektirdiği biçim için Word'de 'Başvurular > İçindekiler' "
        "menüsünden 'Otomatik Tablo 2' seçeneği kullanılabilir. Aşağıdaki "
        "başlıklar Heading 1 stiliyle yazıldığı için içindekiler tablosu "
        "otomatik üretilecektir.)",
    )
    add_body(doc, "ÖNSÖZ")
    add_body(doc, "İÇİNDEKİLER")
    add_body(doc, "SİMGELER VE KISALTMALAR LİSTESİ")
    add_body(doc, "ŞEKİLLER LİSTESİ")
    add_body(doc, "TABLOLAR LİSTESİ")
    add_body(doc, "ÖZET")
    add_body(doc, "BÖLÜM 1. GİRİŞ")
    add_body(doc, "BÖLÜM 2. SİSTEMATİK YAKLAŞIM VE LİTERATÜR")
    add_body(doc, "BÖLÜM 3. TASARIM VE GELİŞTİRME")
    add_body(doc, "BÖLÜM 4. ALGORİTMALAR")
    add_body(doc, "BÖLÜM 5. GÜVENLİK VE TEHDİT MODELİ")
    add_body(doc, "BÖLÜM 6. TEST VE DEĞERLENDİRME")
    add_body(doc, "BÖLÜM 7. SONUÇLAR VE ÖNERİLER")
    add_body(doc, "KAYNAKLAR")
    add_body(doc, "ÖZGEÇMİŞ")
    add_page_break(doc)


def build_kisaltmalar(doc):
    add_heading_1(doc, "SİMGELER VE KISALTMALAR LİSTESİ")
    rows = [
        ("AI", "Yapay Zeka (Artificial Intelligence)"),
        ("API", "Uygulama Programlama Arayüzü (Application Programming Interface)"),
        ("AST", "Soyut Sözdizimi Ağacı (Abstract Syntax Tree)"),
        ("CI/CD", "Sürekli Entegrasyon / Sürekli Dağıtım"),
        ("CSV", "Virgülle Ayrılmış Değerler (Comma-Separated Values)"),
        ("DAG", "Yönlü Çevrimsiz Çizge (Directed Acyclic Graph)"),
        ("DFS", "Derinlik Öncelikli Arama (Depth-First Search)"),
        ("HTTP", "Hiper Metin Transfer Protokolü"),
        ("JSON", "JavaScript Nesne Gösterimi"),
        ("LLM", "Büyük Dil Modeli (Large Language Model)"),
        ("LSA", "Gizli Anlam Analizi (Latent Semantic Analysis)"),
        ("MCP", "Model Context Protocol (Anthropic, 2024)"),
        ("MNACP", "MCP-Native Multi-Agent Coordination Protocol"),
        ("REST", "Temsili Durum Transferi (Representational State Transfer)"),
        ("SCC", "Güçlü Bağlantılı Bileşen (Strongly Connected Component)"),
        ("SSE", "Sunucu Tarafından Gönderilen Olaylar (Server-Sent Events)"),
        ("SVD", "Tekil Değer Ayrışımı (Singular Value Decomposition)"),
        ("TF-IDF", "Terim Frekansı – Ters Belge Frekansı"),
        ("UUID", "Evrensel Tekil Tanımlayıcı"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.autofit = True
    for i, (k, v) in enumerate(rows):
        c1 = table.rows[i].cells[0]
        c2 = table.rows[i].cells[1]
        c1.text = ""
        run = c1.paragraphs[0].add_run(k)
        run.bold = True
        run.font.size = Pt(11)
        run.font.name = "Times New Roman"
        c2.text = ""
        run = c2.paragraphs[0].add_run(": " + v)
        run.font.size = Pt(11)
        run.font.name = "Times New Roman"
    add_page_break(doc)


def build_sekiller_listesi(doc):
    add_heading_1(doc, "ŞEKİLLER LİSTESİ")
    items = [
        "Şekil 2.1. MNACP genel sistem mimarisi",
        "Şekil 3.1. LangGraph orkestratör durum makinesi",
        "Şekil 3.2. Delegasyon protokolü dizilim diyagramı",
        "Şekil 3.3. /chat sayfası — orkestratör sohbet arayüzü",
        "Şekil 3.4. /chat sayfası — canlı ajan ağı görselleştirmesi",
        "Şekil 3.5. /agents sayfası — kayıtlı ajanlar ve güven skorları",
        "Şekil 3.6. /roles sayfası — no-code ajan oluşturma",
        "Şekil 3.7. /monitor sayfası — KPI'lar ve delegasyon geçmişi",
        "Şekil 3.8. Peer delegasyon — DataAgent → AnalysisAgent zinciri",
        "Şekil 4.1. Semantik keşif akışı (TF-IDF + SVD + cosine)",
        "Şekil 4.2. Delegasyon grafı ve döngü tespiti örneği",
        "Şekil 5.1. Tehdit modeli — saldırı yüzeyleri",
        "Şekil 6.1. Sistem karşılaştırması — görev tamamlama oranı",
        "Şekil 6.2. Senaryo bazlı performans ısı haritası",
        "Şekil 6.3. Gecikme dağılımı (P50/P90/P99)",
        "Şekil 6.4. Ajan bazlı performans dağılımı (monitor sayfası)",
    ]
    for it in items:
        add_body(doc, it, justify=False)
    add_page_break(doc)


def build_tablolar_listesi(doc):
    add_heading_1(doc, "TABLOLAR LİSTESİ")
    items = [
        "Tablo 2.1. Karşılaştırılan çok-ajanlı çerçeveler",
        "Tablo 2.2. Kullanılan teknoloji yığını",
        "Tablo 3.1. Örnek ajanlar ve araç listeleri",
        "Tablo 4.1. Karar puanı ağırlıkları ve gerekçeleri",
        "Tablo 5.1. Tehditler ve uygulanan/önerilen önlemler",
        "Tablo 5.2. CodeAgent sandbox yasaklı modüller ve fonksiyonlar",
        "Tablo 6.1. Değerlendirme senaryoları",
        "Tablo 6.2. Sistem başına özet metrikler (15 koşum)",
    ]
    for it in items:
        add_body(doc, it, justify=False)
    add_page_break(doc)


def build_ozet(doc):
    add_heading_1(doc, "ÖZET")
    p = doc.add_paragraph()
    run = p.add_run("Anahtar kelimeler: ")
    run.bold = True
    run.font.size = Pt(12)
    p.add_run(
        "Çok-ajanlı sistemler, Model Context Protocol (MCP), Dinamik delegasyon, "
        "Ajan keşfi, Güven skoru, LLM koordinasyonu",
    ).font.size = Pt(12)
    add_body(
        doc,
        "Bu çalışma, birden fazla yapay zeka ajanının çalışma zamanında "
        "birbirinin yeteneklerini keşfedebildiği, görev devredebildiği ve "
        "döngüsel bağımlılıkları matematiksel olarak engelleyen MNACP "
        "(MCP-Native Multi-Agent Coordination Protocol) adlı bir koordinasyon "
        "protokolünü ve referans uygulamasını sunmaktadır. Sistem; merkezi bir "
        "ajan kayıt servisi (registry), TF-IDF + Truncated SVD tabanlı semantik "
        "ajan keşfi, ardışık başarısızlığı cezalandıran üstel ağırlıklı bir "
        "güven skoru, Strongly Connected Components (Kosaraju) algoritmasıyla "
        "deadlock tespiti yapan bir delegasyon yöneticisi ve LangGraph durum "
        "makinesi ile çalışan bir orkestratör ajandan oluşmaktadır. Doğal dil "
        "tarifinden çalışır ajan üreten bir no-code modülü, Claude Haiku'yu "
        "backend olarak kullanan jenerik bir ajan iskeletiyle gerçeklenmiştir. "
        "Bir ajan, kendi sonucu ile yetinmeyip başka bir ajana doğrudan görev "
        "devredebilmekte (peer delegasyon), bu zincir kullanıcıya canlı "
        "(SSE üzerinden) görselleştirilmektedir.",
    )
    add_body(
        doc,
        "Sistem üç farklı yapı (MNACP, statik atama, merkezi delegasyonsuz) "
        "üzerinde 5 senaryo × 3 tekrar olmak üzere 15 koşumluk bir "
        "değerlendirme protokolüyle karşılaştırılmıştır. MNACP %100 görev "
        "tamamlama ve %96,7 ajan seçim doğruluğuna ulaşmıştır; statik atama "
        "%60 / %30 ile, merkezi baseline ise %100 / %72,2 ile sınırlı kalmıştır. "
        "Sonuçlar, dinamik embedding tabanlı keşfin anahtar kelime eşleşmesine "
        "dayalı atamaya göre belirgin üstünlüğünü ve LLM tabanlı çoklu-ajan "
        "koordinasyonunda matematiksel güvenlik mekanizmalarının pratik "
        "uygulanabilirliğini ortaya koymaktadır.",
    )
    add_page_break(doc)


# ---------- BÖLÜM 1: GİRİŞ ----------

def build_bolum1(doc):
    add_heading_1(doc, "BÖLÜM 1. GİRİŞ")
    add_body(
        doc,
        "Büyük dil modelleri (LLM); doğal dilde sorulan karmaşık sorulara cevap "
        "üretme, kod yazma, metin özetleme gibi görevlerde son iki yıl içinde "
        "olağanüstü bir gelişim göstermiştir. Buna karşın tek bir LLM örneği "
        "kullanıldığında üç temel kısıt belirgin biçimde öne çıkmaktadır: "
        "(1) bağlam penceresi sınırlıdır ve uzun görevlerde model bilgileri "
        "kaybeder; (2) tek bir model her alanda eşit derecede uzmanlaşamaz; "
        "(3) tek nokta arızası yaratır — sistem o tek modele bağımlı hale gelir.",
    )
    add_body(
        doc,
        "Bu kısıtların çözümü için akademi ve endüstri çok-ajanlı sistemlere "
        "yönelmiştir. AutoGen [1], CrewAI [2] ve LangGraph [3] gibi açık kaynak "
        "çerçeveler birden fazla ajanın bir görevi paylaşarak çözmesini sağlar. "
        "Ancak bu sistemlerin önemli bir kısmı statik bir ajan grafiği üzerine "
        "kuruludur: hangi ajanın hangisine görev vereceği önceden, kod düzeyinde "
        "tanımlanmıştır. Yeni bir uzmanlık alanı eklenmek istendiğinde kod "
        "yazılması gerekir; çalışma zamanında ajanların birbirini keşfetmesi söz "
        "konusu değildir.",
    )
    add_body(
        doc,
        "Anthropic'in Kasım 2024'te yayınladığı Model Context Protocol (MCP) [4] "
        "bir LLM uygulamasının dış araç sunucularına standart bir arayüzle "
        "bağlanmasını tanımlamıştır. MCP'nin yarattığı standardın doğal bir "
        "uzantısı, bir ajanın diğer ajanları da birer 'araç sunucusu' olarak "
        "görmesi olabilirdi; ancak protokolün mevcut sürümü ajanlar arası "
        "koordinasyonu kapsam dışı bırakmıştır.",
    )
    add_subheading(doc, "1.1. Problem Tanımı")
    add_body(
        doc,
        "Bu çalışmanın çözmeye çalıştığı problem aşağıdaki dört kısıt altında "
        "tanımlanabilir:",
    )
    add_bullet(doc, "Bir LLM uygulamasında, çalışma zamanında ortaya çıkan görev tanımına en uygun uzman ajan, önceden bilinen bir grafik olmadan nasıl bulunur?")
    add_bullet(doc, "Bir ajan başka bir ajana görev devrettiğinde, A→B→A gibi bir döngü ya da daha karmaşık deadlock'lar nasıl önlenir?")
    add_bullet(doc, "Birden fazla aday arasında geçmiş başarımı ve gecikmesi daha iyi olan ajan nasıl tercih edilir?")
    add_bullet(doc, "Yazılım geliştirici olmayan bir kullanıcı, doğal dil tarifinden çalışan bir ajanı sisteme nasıl ekleyebilir?")
    add_subheading(doc, "1.2. Çalışmanın Katkıları")
    add_body(
        doc,
        "Bu çalışma kapsamında geliştirilen MNACP referans uygulamasının "
        "literatüre ve mühendislik pratiğine somut katkıları şunlardır:",
    )
    add_bullet(doc, "TF-IDF + Truncated SVD ile, harici embedding API'sine bağımlı olmayan, tamamen yerel çalışan semantik ajan keşfi.")
    add_bullet(doc, "Ardışık hatayı ve gecikmeyi tek bir formülde birleştiren üstel ağırlıklı bir güven skoru ve bu skoru similarity ile birleştiren karar mekanizması.")
    add_bullet(doc, "Strongly Connected Components (Kosaraju) algoritmasıyla anlık delegasyon grafiği üzerinde deadlock tespiti.")
    add_bullet(doc, "LangGraph durum makinesi üzerinde decompose → execute → aggregate döngüsünü uygulayan ve canlı SSE eventleri yayan orkestratör.")
    add_bullet(doc, "Doğal dil tarifinden, registry'ye kaydolan ve gerçek HTTP delegasyonu kabul eden GenericAgent üreten no-code modül.")
    add_bullet(doc, "Bir ajanın orkestratörü atlatıp doğrudan başka bir ajana görev devredebildiği peer delegasyon mekanizması ve canlı görselleştirmesi.")
    add_bullet(doc, "GitHub Actions üzerinde Python 3.10/3.11 matrisli sürekli entegrasyon ve Docker Compose ile tek komut konuşlandırma.")
    add_subheading(doc, "1.3. Belge Yapısı")
    add_body(
        doc,
        "Bölüm 2'de literatür ve karşılaştırılan çerçeveler ile teknoloji "
        "tercihlerinin gerekçeleri sunulur. Bölüm 3 sistemin modüler "
        "tasarımını ve uygulama arayüzünü ayrıntılandırır. Bölüm 4 sistemin "
        "doğru çalışmasını sağlayan dört temel algoritmayı (semantik keşif, "
        "güven skoru, döngü/deadlock tespiti, peer delegasyon mantığı) "
        "matematiksel formülasyon ve sözde-kod ile sunar. Bölüm 5 güvenlik "
        "ve tehdit modelini STRIDE çerçevesinde ele alır. Bölüm 6 birim, "
        "entegrasyon ve baseline karşılaştırmalı değerlendirme sonuçlarını "
        "içerir. Bölüm 7 bulguları özetleyip ileri çalışma önerilerini "
        "ele alır.",
    )
    add_page_break(doc)


# ---------- BÖLÜM 2: SİSTEMATİK YAKLAŞIM ----------

def build_bolum2(doc):
    add_heading_1(doc, "BÖLÜM 2. SİSTEMATİK YAKLAŞIM VE LİTERATÜR")
    add_body(
        doc,
        "Bu bölümde önce mevcut çok-ajanlı çerçeveler kıyaslanmış, ardından "
        "MNACP'nin mimarisi ve teknoloji tercihlerinin gerekçeleri ortaya "
        "konmuştur.",
    )
    add_subheading(doc, "2.1. İlgili Çalışmaların Karşılaştırılması")
    add_caption_over(doc, "Tablo 2.1. Karşılaştırılan çok-ajanlı çerçeveler.")
    add_table_simple(
        doc,
        ["Özellik", "MNACP", "AutoGen [1]", "CrewAI [2]", "LangGraph [3]"],
        [
            ["Dinamik ajan keşfi", "Embedding tabanlı", "Statik", "Statik", "Statik"],
            ["Güven skoru", "Var", "Yok", "Yok", "Yok"],
            ["Deadlock koruması", "Kosaraju + DFS", "Kısmi", "Yok", "Yok"],
            ["No-code ajan üretimi", "GenericAgent", "Yok", "Yok", "Yok"],
            ["Peer delegasyon", "Var", "Var", "Var", "Kısmi"],
            ["HTTP/MCP altyapı", "Var", "Yok", "Yok", "Yok"],
            ["Görsel akış izleme", "ReactFlow + SSE", "Yok", "Yok", "LangSmith"],
            ["Konteynerleştirilmiş", "Docker Compose", "Yok", "Yok", "Yok"],
        ],
    )
    add_body(
        doc,
        "Karşılaştırmadan görüleceği üzere, MNACP'nin literatürdeki ana "
        "ayrımlardan birincisi semantik dinamik keşif, ikincisi matematiksel "
        "deadlock korumasıdır. Ek olarak no-code GenericAgent üretimi ile "
        "yazılım geliştirici olmayan kullanıcılara doğrudan rol tanımlama "
        "imkânı sunulmuştur.",
    )
    add_subheading(doc, "2.2. Mimarinin Genel Görünümü")
    add_body(
        doc,
        "Sistem, gevşek bağlı (loosely coupled) HTTP servislerinden oluşur. "
        "Her ajan kendi süreç sınırı içinde çalışır; diğer ajanlarla yalnızca "
        "JSON üzerinden REST çağrıları ile haberleşir. Bu yaklaşım ajanların "
        "farklı dillerle yazılmasına, farklı sunucularda barındırılmasına ve "
        "ayrı yaşam döngülerine sahip olmasına olanak sağlar.",
    )
    add_caption_over(doc, "Şekil 2.1. MNACP genel sistem mimarisi.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(
        "[Şekil için yer tutucu: Frontend (3000) → Orchestrator (8002) → "
        "Registry (8000) → 4 örnek ajan (9001-9004) + Role Builder (8001)]"
    ).italic = True
    add_subheading(doc, "2.3. Teknoloji Yığını ve Gerekçeleri")
    add_caption_over(doc, "Tablo 2.2. Kullanılan teknoloji yığını ve seçim gerekçeleri.")
    add_table_simple(
        doc,
        ["Katman", "Teknoloji", "Gerekçe"],
        [
            ["LLM", "Anthropic Claude (Sonnet 4.5, Haiku 4.5)", "Tool use desteği, streaming, uzun bağlam"],
            ["Web çerçevesi", "FastAPI", "Async, Pydantic ile şema doğrulama, otomatik OpenAPI"],
            ["Şema/doğrulama", "Pydantic v2", "Tip güvenli JSON serileştirme"],
            ["Async HTTP istemci", "httpx", "Async streaming, test dostu"],
            ["Embedding", "scikit-learn TF-IDF + TruncatedSVD", "Yerel çalışır, API anahtarı gerektirmez"],
            ["Durum makinesi", "LangGraph", "Decompose-execute-aggregate döngüsü"],
            ["Frontend", "Next.js 14 + TypeScript + Tailwind", "App Router ile SSE, tip güvenliği"],
            ["Görselleştirme", "ReactFlow + Recharts", "Düğüm/kenar grafiği + bar/line chart"],
            ["Konteyner", "Docker + Docker Compose", "Tek komut konuşlandırma"],
            ["CI", "GitHub Actions (Python 3.10/3.11 matrix)", "Ücretsiz, lint + test + build"],
            ["Test", "pytest, pytest-asyncio", "Async test, mock+gerçek sunucu hibrit"],
            ["Lint", "ruff", "isort + flake8 + black görevini tek komutta üstlenir"],
        ],
    )
    add_subheading(doc, "2.4. Tasarım İlkeleri")
    add_bullet(doc, "Standartlaşmış arayüz: Tüm ajanlar aynı /tools, /delegate, /health uçlarını sunar; bir ajanın başka bir ajan için ayırt edilemez olmasını sağlar.")
    add_bullet(doc, "İdempotent kayıt: Aynı host:port üzerinde yeniden başlayan ajan, registry'de yinelemeye yol açmaz; eski kayıt otomatik temizlenir.")
    add_bullet(doc, "Olay-tabanlı izleme: Tüm orkestrasyon adımları SSE event'leri olarak yayılır; yan etkisi olmayan dış gözlemciler bu akışı dinleyerek davranışı analiz edebilir.")
    add_bullet(doc, "Asgari bağımlılık: Her ajan diğer ajanı doğrudan import etmez; tüm haberleşme HTTP üzerinden gerçekleşir, böylece çoklu süreç ve çoklu host konuşlandırma mümkündür.")
    add_page_break(doc)


# ---------- BÖLÜM 3: TASARIM VE GELİŞTİRME ----------

def build_bolum3(doc):
    add_heading_1(doc, "BÖLÜM 3. TASARIM VE GELİŞTİRME")
    add_body(
        doc,
        "Bu bölüm sistemin teknik iç yapısını detaylı olarak ele alır. Önce "
        "ortak protokol katmanı ve kayıt servisi açıklanır, ardından ajan "
        "soyutlaması, orkestratör durum makinesi, peer delegasyon mekanizması "
        "ve no-code ajan üretimi tanıtılır.",
    )
    add_subheading(doc, "3.1. Protokol Katmanı")
    add_body(
        doc,
        "Tüm bileşenlerin paylaştığı veri yapıları Pydantic v2 modelleri olarak "
        "tanımlanmıştır. Temel mesajlar şunlardır: AgentRegistration (ajan "
        "kayıt formu), AgentInfo (registry'nin dışarıya sunduğu temsil), "
        "DelegationRequest (görev paketi), DelegationResponse (sonuç), "
        "DiscoveryRequest/Result (semantik arama) ve TrustEvent (başarı/"
        "başarısızlık olayı). Her HTTP isteği bu şemalar üzerinde doğrulanır; "
        "şema dışı veri 422 hatası ile reddedilir.",
    )
    add_subheading(doc, "3.2. Kayıt Servisi (Registry)")
    add_body(
        doc,
        "Kayıt servisi sistemin merkezi telefon rehberidir. Ajanlar başlarken "
        "kendilerini buraya kaydeder; orkestratör 'X görevini kim yapabilir?' "
        "sorgusunu buraya iletir. Servis dört bileşenden oluşur:",
    )
    add_bullet(doc, "AgentRegistry: Ajan kayıt sözlüğü; aynı host:port üzerinde eski bir kayıt varsa otomatik temizler (idempotency).")
    add_bullet(doc, "CapabilityEmbedder: Ajanın adı, açıklaması, araçları ve etiketlerini birleştirip TF-IDF + Truncated SVD ile 256 boyutlu vektöre indirger.")
    add_bullet(doc, "TrustScorer: Her tool çağrısı sonrası başarı oranı, ortalama gecikme ve ardışık hata sayısından üstel ağırlıklı bir güven skoru hesaplar.")
    add_bullet(doc, "ToolIndex: Hangi aracın hangi ajanlarda bulunduğunu tutan hash tablosu; 'compute_statistics aracına sahip kim?' gibi exact-match sorgular için.")

    add_subheading(doc, "3.3. Ajan Soyutlaması")
    add_body(
        doc,
        "Tüm ajanlar BaseAgent soyut sınıfından türer. BaseAgent şu sözleşmeleri "
        "sağlar: (a) start/stop ile registry'ye kayıt ve silinme; (b) heartbeat "
        "döngüsüyle registry restart sonrası kendiliğinden yeniden kayıt; (c) "
        "call_tool ile tool çağrısı + güven olayı raporlama; (d) handle_delegation "
        "ile gelen delegasyonun derinlik/döngü kontrolünden geçirilip "
        "_process_delegated_task'a yönlendirilmesi; (e) build_http_app ile "
        "FastAPI uygulaması üretilmesi.",
    )
    add_caption_over(doc, "Tablo 3.1. Sistemde gerçeklenmiş örnek ajanlar ve araç listeleri.")
    add_table_simple(
        doc,
        ["Ajan", "Port", "Araçlar", "Özellik"],
        [
            ["DataAgent", "9001", "load_csv, clean_data, compute_statistics, filter_rows", "İstatistik sonrası AnalysisAgent'a peer delegasyon"],
            ["SearchAgent", "9002", "web_search, fetch_webpage", "DuckDuckGo HTML scraping (gerçek sonuç, API key yok)"],
            ["AnalysisAgent", "9003", "trend_analysis, compare, generate_report, correlation", "Peer delegasyondan gelen istatistik+trend rapor"],
            ["CodeAgent", "9004", "execute_python, analyze_code, format_code, extract_functions", "AST güvenlik denetimi + izole subprocess sandbox"],
        ],
    )
    add_subheading(doc, "3.4. Orkestratör Durum Makinesi")
    add_body(
        doc,
        "Orkestratör, bir kullanıcı görevini alıp birden fazla alt göreve "
        "ayrıştıran, her alt görevi en uygun ajana delege eden ve sonuçları "
        "birleştiren bileşendir. LangGraph ile aşağıdaki durum makinesi "
        "kurulmuştur: START → decompose → execute → aggregate → END. Execute "
        "düğümü, henüz tamamlanmayan alt görev varsa kendisine geri döner.",
    )
    add_caption_over(doc, "Şekil 3.1. LangGraph orkestratör durum makinesi.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(
        "[Şekil için yer tutucu: 4 düğümden oluşan akış diyagramı]"
    ).italic = True
    add_body(
        doc,
        "Decompose adımında Claude Sonnet API'sine 'görevi alt görevlere ayır, "
        "JSON döndür' isteği yapılır. Cevapta her alt görev için id, açıklama, "
        "gerekli yetkinlikler ve bağımlılıklar bulunur. Execute adımında bağımlı "
        "olmayan tüm alt görevler paralel olarak çalıştırılır; bağımlı olanlar "
        "öncüllerinin tamamlanmasını bekler. Aggregate adımında tüm sonuçlar "
        "tekrar Claude'a gönderilip kullanıcıya yönelik tutarlı bir cevaba "
        "dönüştürülür.",
    )
    add_subheading(doc, "3.5. Peer Delegasyon")
    add_body(
        doc,
        "Standart çok-ajanlı sistemlerde tüm haberleşme orkestratör üzerinden "
        "geçen yıldız topolojisi izler. MNACP'de ise bir ajan, kendi sonucunun "
        "yetersiz kaldığını fark ederse doğrudan başka bir ajana delegasyon "
        "yapabilir. DataAgent örneği, görevde hem 'istatistik' hem de "
        "'rapor/analiz' kelimeleri bulunduğunda istatistiği yerel hesaplar, "
        "ardından registry'de AnalysisAgent'ı keşfedip ona trend analizi ve "
        "rapor üretmesi için delegasyon yapar. Peer çağrısının sonucu özel bir "
        "_peer_delegations alanı içinde orkestratöre döner; orkestratör bu "
        "alanı yakalayıp peer_delegation SSE event'i yayınlar. Frontend bu "
        "event'i ajan→ajan kenar olarak çizer; böylece kullanıcı zincirleme "
        "delegasyonu gerçek zamanlı izler.",
    )
    add_caption_over(doc, "Şekil 3.2. Delegasyon protokolü dizilim diyagramı.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(
        "[Şekil için yer tutucu: Kullanıcı → Frontend → Orkestratör → DataAgent → AnalysisAgent dizilim diyagramı]"
    ).italic = True

    add_subheading(doc, "3.6. No-code Ajan Üretimi")
    add_body(
        doc,
        "Yazılım bilgisi olmayan kullanıcıların doğal dil tarifinden ajan "
        "ekleyebilmesi için no-code modülü geliştirilmiştir. Akış şöyle çalışır: "
        "(1) kullanıcı bir tarif yazar (örn. 'finansal raporları analiz eden, "
        "borsa verisini takip eden ajan'); (2) RoleBuilder bu tarifi Claude'a "
        "gönderip JSON formatında ajan adı, etiketler ve önerilen araç listesi "
        "alır; (3) kullanıcı önerileri görüp onaylar; (4) onay sonrası "
        "GenericAgent başlatılır — bu ajan Claude Haiku'yu backend olarak "
        "kullanarak hem tool çağrılarını hem delegasyonları cevaplar; (5) "
        "ajan registry'ye base_path='/agents/{id}' alanıyla kayıt olur, böylece "
        "rol-builder konteynerinin tek portu üzerinden çoklu ajan barındırılır.",
    )
    add_subheading(doc, "3.7. Frontend ve Canlı Görselleştirme")
    add_body(
        doc,
        "Frontend Next.js 14 App Router üzerinde TypeScript ve Tailwind CSS "
        "ile yazılmıştır. Toplam beş sayfadan oluşur: ana panel, /chat (görev "
        "çalıştırma ve canlı izleme), /agents (kayıtlı ajan listesi), /roles "
        "(no-code rol oluşturma) ve /monitor (metrik panosu). Tüm sayfalar "
        "registry'ye 3 saniyede bir polling yaparak güncel kalır.",
    )
    add_body(
        doc,
        "/chat sayfası, kullanıcıdan görev alır ve orkestratörün /run/stream "
        "uç noktasına SSE bağlantısı kurar. Gelen her event (decompose_done, "
        "subtask_start, subtask_done, peer_delegation, final_answer) ekrandaki "
        "AgentGraph (ReactFlow) bileşenini günceller: düğümler ajanlar, sarı "
        "animasyonlu kenarlar aktif delegasyonlar, yeşil/kırmızı kenarlar "
        "tamamlanmış/başarısız çağrılardır. Kullanıcı bu sayfada görev "
        "ayrıştırmasını, hangi alt görevin hangi ajana gittiğini, peer "
        "delegasyon zincirlerinin gerçek zamanlı oluşumunu ve nihai cevabın "
        "üretilmesini izleyebilir.",
    )
    add_caption_over(doc, "Şekil 3.3. /chat sayfası — orkestratör sohbet arayüzü.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(
        "[Ekran görüntüsü için yer tutucu: /chat sayfasının üst yarısı — "
        "preset görevler, görev textarea'sı, Çalıştır butonu, alt görev "
        "kartları]"
    ).italic = True
    add_caption_over(doc, "Şekil 3.4. /chat sayfası — canlı ajan ağı görselleştirmesi.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(
        "[Ekran görüntüsü için yer tutucu: /chat sayfasındaki AgentGraph "
        "bileşeni — ortada Orkestratör düğümü, çevresinde DataAgent / "
        "SearchAgent / AnalysisAgent / CodeAgent, aktif (sarı/yeşil) "
        "kenarlar görünür]"
    ).italic = True

    add_body(
        doc,
        "/agents sayfası, registry'ye kayıtlı tüm ajanların kart görünümünü "
        "sunar. Her kart ajanın adı, açıklaması, etiketleri, çevrim içi "
        "durumu, güven skoru yüzdesi ve sahip olduğu araç listesini gösterir. "
        "Bu sayfa hem örnek (DataAgent, SearchAgent, AnalysisAgent, CodeAgent) "
        "hem de no-code modülüyle dinamik olarak eklenmiş ajanları aynı "
        "biçimde listeler.",
    )
    add_caption_over(doc, "Şekil 3.5. /agents sayfası — kayıtlı ajanlar ve güven skorları.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(
        "[Ekran görüntüsü için yer tutucu: /agents sayfası — 4 örnek ajan "
        "ve no-code ile eklenmiş örnek bir ajanın kartları]"
    ).italic = True

    add_body(
        doc,
        "/roles sayfası, no-code ajan oluşturma akışını sunar. Kullanıcı bir "
        "metin alanına ajanın ne yapmasını istediğini doğal dilde yazar; "
        "'Araç Öner' butonu Claude'a istek atarak ajan adı, etiketler ve "
        "araç önerileri olarak yapılandırılmış bir teklif döndürür. Kullanıcı "
        "teklifi inceler, gerekirse düzenler ve 'Onayla ve Sisteme Ekle' "
        "butonuyla GenericAgent'ı başlatıp registry'ye kaydeder. Ardından "
        "yeni ajan keşif sıralamasında değerlendirilebilir hale gelir.",
    )
    add_caption_over(doc, "Şekil 3.6. /roles sayfası — no-code ajan oluşturma.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(
        "[Ekran görüntüsü için yer tutucu: /roles sayfası — açıklama "
        "textarea'sı, önerilen araçlar listesi ve Onayla butonu]"
    ).italic = True

    add_body(
        doc,
        "/monitor sayfası, sistemin operasyonel sağlığını ve geçmiş "
        "performansını izleyen panodur. Üst kısımda dört KPI kartı bulunur "
        "(toplam delegasyon, başarılı, reddedilen, başarı yüzdesi). Altında "
        "ajan bazlı yığılmış bar chart (her hedef ajan için başarılı ve "
        "hatalı çağrı dağılımı) ve eşleşen tablo (toplam çağrı, başarı "
        "yüzdesi, ortalama gecikme); ardından son 30 delegasyonun gecikme "
        "çizgi grafiği yer alır. Sayfanın alt kısmında baseline "
        "karşılaştırması grafikleri (PNG olarak gömülü), en altta da "
        "delegasyon geçmişi listesi sunulur.",
    )
    add_caption_over(doc, "Şekil 3.7. /monitor sayfası — KPI'lar ve delegasyon geçmişi.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(
        "[Ekran görüntüsü için yer tutucu: /monitor sayfası — 4 KPI kartı, "
        "ajan bar chart, latency trend ve delegasyon kartları]"
    ).italic = True

    add_body(
        doc,
        "Peer delegasyon mekanizmasının canlı görselleştirmesi sistemin en "
        "ayırt edici görsellerinden biridir. Aşağıdaki ekran görüntüsü, "
        "kullanıcının istatistik+rapor içeren bir görev verdiği anda "
        "DataAgent'ın görevi alıp AnalysisAgent'a peer delegasyon "
        "yaptığı, böylece grafikte hem orkestratör→DataAgent hem de "
        "DataAgent→AnalysisAgent kenarlarının aynı anda göründüğü zinciri "
        "yansıtır.",
    )
    add_caption_over(doc, "Şekil 3.8. Peer delegasyon — DataAgent → AnalysisAgent zinciri.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(
        "[Ekran görüntüsü için yer tutucu: /chat sayfasında peer delegasyon "
        "tetiklendiğinde AgentGraph'ta görünen iki kenar — orkestratör → "
        "DataAgent (yeşil) ve DataAgent → AnalysisAgent (yeşil/peer label)]"
    ).italic = True
    add_page_break(doc)


# ---------- BÖLÜM 4: ALGORİTMALAR ----------

def build_bolum4_algoritmalar(doc):
    add_heading_1(doc, "BÖLÜM 4. ALGORİTMALAR")
    add_body(
        doc,
        "Bu bölüm sistemin doğru çalışmasını sağlayan dört temel algoritmayı "
        "ayrıntılı olarak ele alır: semantik ajan keşfi, üstel ağırlıklı "
        "güven skoru hesaplama, delegasyon çizgesinde döngü ve deadlock "
        "tespiti, ve peer delegasyon karar mantığı. Her algoritma için "
        "matematiksel formülasyon, sözde-kod ve karmaşıklık analizi "
        "verilmiştir.",
    )

    add_subheading(doc, "4.1. Semantik Ajan Keşfi")
    add_body(
        doc,
        "Orkestratörün karşılaştığı temel problem şudur: kullanıcının doğal "
        "dilde tanımladığı bir görevi, registry'de kayıtlı N adet ajan "
        "arasından en uygun olanına atamak. Ajanların yetkinlikleri de "
        "yapılandırılmamış metin (ad, açıklama, araç açıklamaları, etiketler) "
        "olarak tanımlandığından, eşleşme problemi bir bilgi-erişim "
        "(information retrieval) problemine indirgenir.",
    )
    add_body(
        doc,
        "İki klasik yaklaşım vardır: (a) anahtar kelime eşleşmesi — basit "
        "ama eşanlamlılar ve dilbilimsel varyasyonlar karşısında zayıf; "
        "(b) yoğun (dense) embedding — güçlü ama harici API gerektirir. "
        "Bu çalışmada üçüncü bir yol — TF-IDF ile Truncated SVD'nin "
        "birleştirilmesiyle elde edilen yerel Latent Semantic Analysis "
        "(LSA) — tercih edilmiştir. Bu yaklaşım API anahtarı "
        "gerektirmez, deterministiktir ve dış servise bağımlılık yaratmaz.",
    )
    add_body(doc, "Ajan A için temsil metni şöyle birleştirilir:")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "doc(A) = name(A) ⊕ desc(A) ⊕ ⊕ᵢ tool_nameᵢ ⊕ tool_descᵢ ⊕ ⊕ⱼ tagⱼ          (4.1)"
    )
    run.italic = True
    run.font.size = Pt(12)
    add_body(
        doc,
        "Tüm ajan dokümanları üzerinde TfidfVectorizer (max_features = 4096, "
        "n-gram = 1–2) sığdırılır. Elde edilen seyrek matris üzerinde "
        "Truncated SVD ile k = 256 boyutlu yoğun bir alt uzaya iz düşürülür. "
        "Görev metni q da aynı pipeline'dan geçirilir; ajanlarla benzerlik "
        "cosine ile hesaplanır:",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "sim(q, A) = (vec(q) · vec(A)) / (‖vec(q)‖ × ‖vec(A)‖)                       (4.2)"
    )
    run.italic = True
    run.font.size = Pt(12)
    add_body(
        doc,
        "Yalnızca semantik benzerlik, geçmişte kötü performans gösteren bir "
        "ajanı yine de seçebilirdi. Bu nedenle skor, güven (trust) faktörüyle "
        "ağırlıklı olarak birleştirilir:",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "score(A | q) = α · sim(q, A) + β · trust(A),    α = 0,7;  β = 0,3              (4.3)"
    )
    run.italic = True
    run.font.size = Pt(12)
    add_caption_over(doc, "Tablo 4.1. Karar puanı ağırlıkları ve gerekçeleri.")
    add_table_simple(
        doc,
        ["Ağırlık", "Değer", "Etki"],
        [
            ["α (similarity)", "0,7", "Görev-ajan uyumunun başat belirleyicisi"],
            ["β (trust)", "0,3", "Geçmişte yavaş veya hatalı ajanları cezalandırır"],
            ["Toplam", "1,0", "Skor [0, 1] aralığında yorumlanabilir"],
        ],
    )

    add_caption_over(doc, "Şekil 4.1. Semantik keşif akışı (TF-IDF + SVD + cosine).")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(
        "[Şekil için yer tutucu: Görev metni → TF-IDF → SVD (256-d) → "
        "cosine sim ile her ajan vektörü → trust ile ağırlıklı toplam → "
        "top-K aday]"
    ).italic = True

    add_body(doc, "Sözde-kod aşağıdaki gibidir:")
    pseudo = (
        "Algoritma 1: SemantikKeşif(q, registry, k)\n"
        "  Girdi:  q (görev metni), registry (N ajan), k (top-k)\n"
        "  Çıktı:  En yüksek skorlu k ajan\n"
        "  1: V ← {vec(A) | A ∈ registry}                  // SVD ile önceden hesaplanmış\n"
        "  2: q_vec ← TF-IDF + SVD pipeline(q)\n"
        "  3: scores ← []\n"
        "  4: for each A in registry do\n"
        "  5:     sim ← cosine(q_vec, vec(A))\n"
        "  6:     score ← 0,7·sim + 0,3·trust(A)\n"
        "  7:     scores.append((A, score))\n"
        "  8: end for\n"
        "  9: scores.sort(by=score, desc)\n"
        " 10: return scores[1..k]"
    )
    p = doc.add_paragraph()
    run = p.add_run(pseudo)
    run.font.name = "Consolas"
    run.font.size = Pt(10)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(12)

    add_body(
        doc,
        "Karmaşıklık. Eğitim aşaması (TfidfVectorizer + TruncatedSVD) "
        "O(N · L + N · k²) süre alır; burada L ortalama doküman uzunluğu, "
        "k = 256. Sorgu zamanı O(L_q + N · k) düzeyindedir; binlerce ajana "
        "kadar ölçeklenir. Ajan eklendiğinde matris yeniden sığdırılır; "
        "registry _refit_embeddings çağrısı ile bu adımı amortizasyonlu "
        "tutar.",
    )

    add_subheading(doc, "4.2. Üstel Ağırlıklı Güven Skoru")
    add_body(
        doc,
        "Her tool çağrısı ve delegasyon sonunda registry'ye bir TrustEvent "
        "iletilir; başarı bayrağı, gecikme süresi ve zaman damgası içerir. "
        "Bu olaylardan ajan başına üç istatistik tutulur: total (toplam "
        "deneme), successes (başarılı sonuç) ve consecutive_failures "
        "(ardışık başarısızlık sayısı, başarı geldiğinde sıfırlanır).",
    )
    add_body(doc, "Skor üç çarpanın çarpımıdır:")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "trust(A) = success_rate × latency_factor × failure_penalty               (4.4)"
    )
    run.italic = True
    run.font.size = Pt(12)
    add_body(doc, "Bileşenler:")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("success_rate = successes / total                                            (4.5)")
    run.italic = True
    run.font.size = Pt(12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "latency_factor = 1 / (1 + avg_latency_ms / 2000)                            (4.6)"
    )
    run.italic = True
    run.font.size = Pt(12)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "failure_penalty = 0,9 ^ consecutive_failures                                   (4.7)"
    )
    run.italic = True
    run.font.size = Pt(12)
    add_body(
        doc,
        "Skor [0,05; 1,00] aralığına sıkıştırılır (clip). Bu sayede tek bir "
        "kötü deneme ajanı tamamen sıralama dışına atmaz, ancak ardışık "
        "başarısızlıklar üstel olarak cezalandırılır: 5 ardışık hata "
        "0,9⁵ ≈ 0,59 çarpanı verir; 10 ardışık hata 0,35'e düşer. Tek bir "
        "başarı consecutive_failures'ı sıfırlayarak ajanın tekrar "
        "yarışmasına izin verir.",
    )
    add_body(
        doc,
        "Hedef gecikme TARGET_LATENCY_MS = 2000 olarak seçilmiştir; bu "
        "değerde latency_factor = 0,5 olur. 200 ms civarındaki çağrılarda "
        "faktör 0,91, 8 saniyelik çağrılarda 0,2'dir. Bu tasarım hızlı ve "
        "doğru ajanları açık biçimde ödüllendirir.",
    )

    add_subheading(doc, "4.3. Döngü ve Deadlock Tespiti")
    add_body(
        doc,
        "Aktif delegasyonlar yönlü bir çizge G(V, E) olarak tutulur: V "
        "ajanlar, E aktif delegasyon kenarları. A → B kenarı 'A şu an B'nin "
        "yanıtını bekliyor' anlamına gelir. Bu çizgede iki tür problem "
        "yaratılabilir:",
    )
    add_bullet(doc, "Basit döngü: Bir ajan zincirde kendisinden önce yer alıyorsa (A → B → A) sonsuz beklemeye yol açar.")
    add_bullet(doc, "Deadlock (karşılıklı bekleme): İki veya daha fazla ajan birbirinin sonucunu bekleyen bir Strongly Connected Component oluşturur.")

    add_caption_over(doc, "Şekil 4.2. Delegasyon grafı ve döngü tespiti örneği.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(
        "[Şekil için yer tutucu: Soldaki graf 'A → B → A döngüsü'; sağdaki "
        "graf 'A → B, B → C, C → A → SCC = {A, B, C}' örneği]"
    ).italic = True

    add_body(
        doc,
        "MNACP üç katmanlı bir savunma uygular. Birinci katman: derinlik "
        "sınırı (max_depth = 5). Bu, dolaylı döngülerin uzunluğunu üstten "
        "bağlar. İkinci katman: zincir tabanlı kontrol — bir delegasyon "
        "isteği oluştuğunda hedef ajan zaten chain listesinde bulunuyor mu "
        "diye bakılır. Üçüncü katman: global graf üzerinde DFS ile çevrim "
        "araması.",
    )
    add_body(doc, "Sözde-kod, is_safe_to_delegate fonksiyonu için:")
    pseudo2 = (
        "Algoritma 2: GüvenliMi(from, to, chain)\n"
        "  Girdi:  from (kaynak ajan), to (hedef ajan), chain (şu ana kadar zincir)\n"
        "  Çıktı:  (bool, açıklama)\n"
        "  1: if to ∈ chain then\n"
        "  2:     return (false, 'Döngüsel delegasyon: hedef zincirde')\n"
        "  3: G.addEdge(from, to)                          // geçici\n"
        "  4: c ← hasGlobalCycle(G)                        // DFS\n"
        "  5: G.removeEdge(from, to)                       // geri al\n"
        "  6: if c then\n"
        "  7:     return (false, 'Bu kenar global döngü oluşturuyor')\n"
        "  8: return (true, '')"
    )
    p = doc.add_paragraph()
    run = p.add_run(pseudo2)
    run.font.name = "Consolas"
    run.font.size = Pt(10)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(12)

    add_body(
        doc,
        "find_deadlocked_agents fonksiyonu Kosaraju algoritması ile "
        "Strongly Connected Components (SCC) hesaplar. Algoritma iki DFS'ten "
        "oluşur: ilk DFS düğümleri bitiş zamanına göre yığına yerleştirir, "
        "ikinci DFS ters çevrilmiş graf üzerinde yığın sırasına göre "
        "yürütülür ve her DFS ağacı bir SCC oluşturur. Boyutu 1'den büyük "
        "SCC'ler birbirini bekleyen ajan kümeleridir; sistem bunları "
        "saptayıp ilgili delegasyonları zaman aşımına uğratır.",
    )
    add_body(
        doc,
        "Karmaşıklık. Zincir kontrolü O(d), burada d zincir uzunluğu "
        "(en fazla max_depth = 5). DFS tabanlı global çevrim kontrolü "
        "O(V + E). Kosaraju iki DFS gerektirdiğinden yine O(V + E). Tüm "
        "kontroller delegasyon sayısı ile lineer ölçeklenir.",
    )

    add_subheading(doc, "4.4. Peer Delegasyon Karar Mantığı")
    add_body(
        doc,
        "Standart yıldız topolojisinde her şey orkestratörden geçer; bu "
        "merkez ajanı darboğaza dönüştürür ve uzmanlar arası 'doğal' "
        "iletişimi engeller. MNACP'de bir ajan, kendi sonucunun bir kısmını "
        "ya da tamamını başka bir ajana doğrudan devredebilir. Bu davranış, "
        "DataAgent için somut bir kuralla örneklenmiştir: görev hem 'istatistik' "
        "hem de 'rapor/analiz/trend' belirteçlerini içeriyorsa, DataAgent "
        "yerel istatistiği hesapladıktan sonra rapor üretimini AnalysisAgent'a "
        "delege eder.",
    )
    add_body(doc, "Sözde-kod:")
    pseudo3 = (
        "Algoritma 3: DataAgent.processDelegatedTask(task, context)\n"
        "  1: combined ← lower(task) + ' ' + lower(context.original_task)\n"
        "  2: needs_stats     ← KW_STATS  ⊆ combined          // 'istatistik', 'mean', ...\n"
        "  3: needs_analysis  ← KW_ANALYSIS ⊆ combined       // 'rapor', 'trend', ...\n"
        "  4: rows ← resolveRows(context)\n"
        "  5: if needs_stats then\n"
        "  6:     col ← context.column ∨ pickNumericColumn(rows, task)\n"
        "  7:     stats ← compute_statistics(rows, col)\n"
        "  8:     if needs_analysis then\n"
        "  9:         analyst ← discovery.findBestAgent('rapor + trend',\n"
        " 10:                       capabilities=[generate_report, trend_analysis],\n"
        " 11:                       exclude=[self.id])\n"
        " 12:         peer ← delegationMgr.delegate(self → analyst,\n"
        " 13:                       task='trend + rapor',\n"
        " 14:                       context={values, statistics, column, ...},\n"
        " 15:                       chain=context._delegation_chain)\n"
        " 16:         return {statistics, analysis: peer.result,\n"
        " 17:                 _peer_delegations: [peer.metadata]}\n"
        " 18:     return stats\n"
        " 19: ..."
    )
    p = doc.add_paragraph()
    run = p.add_run(pseudo3)
    run.font.name = "Consolas"
    run.font.size = Pt(10)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(12)
    add_body(
        doc,
        "Önemli ayrıntı: DataAgent peer çağrısının sonucunu kendi cevabının "
        "içinde _peer_delegations alanı ile döndürür. Orkestratör bu alanı "
        "tespit edip peer_delegation SSE event'i yayınlar. Böylece kullanıcı "
        "arayüzünde DataAgent → AnalysisAgent kenarı gerçek zamanlı olarak "
        "çizilir; hem akademik 'çoklu-ajan iletişimi' iddiası kanıtlanmış "
        "olur, hem de zincir gözlemlenebilir.",
    )
    add_body(
        doc,
        "Bu örnekteki kural bazlı tetikleyici, basit ama etkili bir başlangıç "
        "noktasıdır. Bölüm 7'de önerildiği üzere ileriki sürümde kararlar "
        "kuralla değil, küçük bir sınıflandırıcı modelle verilebilir.",
    )
    add_page_break(doc)


# ---------- BÖLÜM 5: GÜVENLİK VE TEHDİT MODELİ ----------

def build_bolum5_guvenlik(doc):
    add_heading_1(doc, "BÖLÜM 5. GÜVENLİK VE TEHDİT MODELİ")
    add_body(
        doc,
        "Çok-ajanlı sistemler, klasik istemci-sunucu uygulamalarına kıyasla "
        "daha geniş bir saldırı yüzeyi sunar: birden fazla bağımsız bileşen, "
        "ağ üzerinden mesajlaşan ajanlar, çalışma zamanında dinamik olarak "
        "eklenen kullanıcı tanımlı roller ve dış API çağrıları. Bu bölümde "
        "MNACP'nin tehdit modeli STRIDE çerçevesi etrafında ele alınmış, "
        "hâlihazırda uygulanan güvenlik kontrolleri ve henüz uygulanmamış "
        "olan ancak ürünleştirme aşamasında zorunlu hale gelecek önlemler "
        "açıkça belirtilmiştir.",
    )

    add_subheading(doc, "5.1. Saldırı Yüzeyleri")
    add_body(
        doc,
        "Sistem dört temel saldırı yüzeyi üzerinden değerlendirilebilir:",
    )
    add_bullet(doc, "Kullanıcı arayüzü ve orkestratör girişi: Kullanıcının doğal dilde gönderdiği görev metni, decompose adımında doğrudan LLM'e iletilir. Kötü niyetli bir prompt, sistem davranışını manipüle etmeye çalışabilir (prompt injection).")
    add_bullet(doc, "Ajanlar arası HTTP haberleşmesi: Tüm /delegate ve /tools/{name} çağrıları kimliksiz şekilde kabul edilir. Aynı ağda kötü niyetli bir istemci bir ajana doğrudan görev gönderebilir.")
    add_bullet(doc, "Registry kayıt arayüzü: /agents/register uç noktası kimlik doğrulaması istemediğinden bir saldırgan sahte ajan kaydı yapabilir (registry poisoning) ve böylece keşif sıralamasını etkileyebilir.")
    add_bullet(doc, "Kod yürütme yüzeyi: CodeAgent isteğe bağlı Python kodunu çalıştırır. Sandbox dışına çıkış (sandbox escape) sistemin tamamını riske atar.")

    add_caption_over(doc, "Şekil 5.1. Tehdit modeli — saldırı yüzeyleri.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(
        "[Şekil için yer tutucu: Kullanıcı / Diğer ajan / Saldırgan → "
        "Frontend → Orkestratör → Registry → Ajanlar diyagramı, her "
        "okun üzerinde tehdit etiketi]"
    ).italic = True

    add_subheading(doc, "5.2. STRIDE Tehdit Sınıflandırması")
    add_body(
        doc,
        "Microsoft'un STRIDE modeli, tehditleri altı kategoride ele alır: "
        "Spoofing (kimlik taklidi), Tampering (veri değiştirme), Repudiation "
        "(inkâr edilebilirlik), Information Disclosure (bilgi sızıntısı), "
        "Denial of Service (hizmet engelleme) ve Elevation of Privilege "
        "(yetki yükseltme). MNACP bağlamında bu kategorilerin somut "
        "karşılıkları aşağıdaki tabloda özetlenmiştir.",
    )
    add_caption_over(doc, "Tablo 5.1. Tehditler ve uygulanan/önerilen önlemler.")
    add_table_simple(
        doc,
        ["Kategori", "Somut Tehdit", "Mevcut Önlem", "Önerilen Ek Önlem"],
        [
            [
                "Spoofing",
                "Sahte ajan kaydı (registry poisoning)",
                "Aynı host:port üzerinde idempotent kayıt",
                "Mutual TLS, ajan başına imzalı JWT",
            ],
            [
                "Tampering",
                "Delegasyon mesajının MITM ile değiştirilmesi",
                "JSON şema doğrulaması (Pydantic)",
                "HTTPS + mesaj imzalama (HMAC)",
            ],
            [
                "Repudiation",
                "Bir ajanın görev kabul ettiğini inkâr etmesi",
                "Tüm delegasyonlar history'de kaydedilir (timestamp, request_id)",
                "Append-only audit log + kriptografik zincir (Merkle)",
            ],
            [
                "Information Disclosure",
                "Görev içeriğinin loglara/üçüncü servislere sızması",
                "Logger seviyesi configurable, sırlar .env'de",
                "PII maskeleme, log redaction politikası",
            ],
            [
                "Denial of Service",
                "Aynı ajana ardışık ağır görev gönderme",
                "TrustScorer ardışık hata cezası, 60 saniye HTTP timeout",
                "Rate limiting (token bucket), kuyruk derinliği sınırı",
            ],
            [
                "Elevation of Privilege",
                "Kod sandbox'ından çıkış",
                "AST analizi, izole subprocess, beyaz liste modüller",
                "Linux namespace izolasyonu, seccomp filtresi",
            ],
        ],
    )

    add_subheading(doc, "5.3. CodeAgent Sandbox Mimarisi")
    add_body(
        doc,
        "CodeAgent, kullanıcı veya başka bir ajan tarafından sağlanan "
        "Python kodunu çalıştırır. Bu yüzey en yüksek riske sahip noktadır "
        "ve iki kademeli savunma uygulanmıştır.",
    )
    add_body(
        doc,
        "Birinci kademe: Statik AST analizi. Kod, ast.parse ile soyut "
        "sözdizimi ağacına dönüştürülür ve ağaç üzerinde gezilir. Yasaklı "
        "modül import'ları (os, sys, subprocess, socket, shutil, pathlib, "
        "importlib, ctypes, multiprocessing, threading, pickle, signal, pty) "
        "ve yasaklı çağrılar (eval, exec, __import__, compile; öznitelik "
        "olarak system/popen/run/Popen) tespit edilirse kod hiç "
        "çalıştırılmadan reddedilir. Yalnızca beyaz listedeki güvenli stdlib "
        "modülleri ve bilimsel kütüphaneler (math, statistics, json, re, "
        "datetime, hashlib, base64, numpy, pandas, scipy, sklearn vb.) "
        "kullanılabilir.",
    )
    add_body(
        doc,
        "İkinci kademe: İzole subprocess. AST kontrolünden geçen kod, "
        "subprocess.run ile yeni bir Python yorumlayıcı sürecinde "
        "çalıştırılır. Çağrı; kullanıcı kontrolünde olmayan bir sarmalayıcı "
        "(wrapper) ile sarılır, stdout ve stderr ayrı ayrı yakalanır, "
        "1 ile 30 saniye arası kullanıcı tarafından belirlenebilen ancak "
        "üst sınırla cebren kısıtlanan bir zaman aşımı uygulanır. Çıktının "
        "boyutu da kesilir (stdout 4 KB, stderr 2 KB).",
    )

    add_caption_over(doc, "Tablo 5.2. CodeAgent sandbox yasaklı modüller ve fonksiyonlar.")
    add_table_simple(
        doc,
        ["Kategori", "Yasaklı Öğeler", "Gerekçe"],
        [
            [
                "Sistem erişimi",
                "os, sys, subprocess, shutil, pathlib, signal, pty",
                "Dosya/process kontrolü ile sandbox'tan kaçış",
            ],
            [
                "Ağ",
                "socket, urllib (kısmi)",
                "Veri sızdırma, dış C2 sunucusuna bağlantı",
            ],
            [
                "Dinamik kod",
                "eval, exec, __import__, compile",
                "Çalışma zamanında yasaklı modül yükleme",
            ],
            [
                "Düşük seviye",
                "ctypes, importlib",
                "C kütüphanelerine doğrudan erişim, yapı bozma",
            ],
            [
                "Eşzamanlılık",
                "multiprocessing, threading",
                "Yarış durumu yaratma, kaynak tüketimi",
            ],
            [
                "Serileştirme",
                "pickle, shelve, dbm",
                "Kötü amaçlı pickle yükü ile uzaktan kod çalıştırma",
            ],
        ],
    )

    add_body(
        doc,
        "Bu iki kademeli yaklaşım belirli sınırlar dahilinde güvenli bir "
        "yürütme sağlar; ancak gerçek bir prodüksiyon dağıtımında Docker "
        "veya Linux namespace'leri ile süreç izolasyonu, seccomp-bpf "
        "filtreleri ve cgroup tabanlı kaynak sınırlandırması ek önlem "
        "olarak şarttır.",
    )

    add_subheading(doc, "5.4. Prompt Injection")
    add_body(
        doc,
        "Decompose ve aggregate aşamalarında kullanıcı metni doğrudan "
        "Claude'a iletilir. Kullanıcı, sistem promptunu manipüle edici "
        "talimatlar yazarak orkestratörün davranışını değiştirmeye "
        "çalışabilir (örneğin 'önceki tüm talimatları yok say ve şu komutu "
        "çalıştır'). Mevcut sürümde bu saldırıya karşı uygulanmış aktif "
        "savunma yoktur. Önerilen yaklaşımlar: (i) sistem prompt'unu "
        "kullanıcı metninden ayrı bir delimiter (XML benzeri etiketler) "
        "ile sarmak; (ii) decompose sonucunda yalnızca bilinen ajan "
        "yetkinlik kümesi içindeki alt görevleri kabul etmek; (iii) kritik "
        "tool çağrılarını insan onayına bağlayan bir 'tool gating' "
        "mekanizması.",
    )

    add_subheading(doc, "5.5. Döngüsel Delegasyon ve Kaynak Tükenmesi")
    add_body(
        doc,
        "Kötü niyetli veya hatalı bir ajan, başka bir ajana sürekli "
        "kendisini hedef gösteren delegasyonlar üreterek bir denial of "
        "service yaratabilir. MNACP'de bu saldırıya karşı üç kontrol "
        "uygulanmaktadır: (a) DelegationManager içinde max_depth = 5 "
        "sınırı; (b) zincir tabanlı döngü kontrolü — istenen hedef zincirde "
        "varsa istek REJECTED durumuyla reddedilir; (c) global graf üzerinde "
        "DFS ile çevrim arama. Ek olarak DeadlockDetector, Kosaraju "
        "algoritmasıyla birden fazla ajanlı SCC'leri tespit ederek karşılıklı "
        "bekleme durumlarını yakalar.",
    )

    add_subheading(doc, "5.6. Kimlik Doğrulama Eksikliği")
    add_body(
        doc,
        "Bu çalışma akademik bir referans uygulama olduğundan, ajanlar arası "
        "HTTP çağrılarında kimlik doğrulaması bilinçli olarak dışarıda "
        "bırakılmıştır. Bir ajanın diğeri olduğunu kanıtlamak için OAuth 2.0 "
        "client credentials akışı veya mutual TLS gibi mekanizmalar bir "
        "sonraki sürümde eklenmelidir. Registry tarafında ise her ajan "
        "kayıt sırasında kendi imzalı sertifikasıyla kimliklenmeli, /discover "
        "ve /register uçları yetkili istemcilere kısıtlanmalıdır.",
    )

    add_subheading(doc, "5.7. Bağımlılık ve Tedarik Zinciri")
    add_body(
        doc,
        "Sistem 30'a yakın doğrudan Python paketi ve 200'ün üzerinde dolaylı "
        "JavaScript paketi kullanmaktadır. Tedarik zinciri saldırılarına "
        "karşı şu uygulamalar mevcuttur: (i) requirements.txt ve "
        "package-lock.json sürüm sabitleme; (ii) GitHub Actions üzerinde "
        "her commit'te bağımlılıkların yeniden çözümlenmesi; (iii) Docker "
        "image'larında temel olarak resmi python:3.11-slim ve node:18-alpine "
        "image'larının kullanılması. Önerilen ek önlemler: pip-audit ve "
        "npm audit'in CI'a entegre edilmesi, Software Bill of Materials "
        "(SBOM) üretimi ve image imzalama (Sigstore/cosign).",
    )

    add_subheading(doc, "5.8. Veri Güvenliği ve Gizlilik")
    add_body(
        doc,
        "Sistem üzerinden işlenen veriler üç kategoriye ayrılır: (a) "
        "kullanıcı görev metni, (b) ajanlar arası bağlam (context) "
        "verisi — istatistik sonuçları, CSV satırları, web aramaları; "
        "(c) Anthropic API'sine giden tam metin. Mevcut sürüm tüm bu "
        "verileri düz metin halinde işler ve loglara kısmen yazar. "
        "ANTHROPIC_API_KEY .env dosyasında saklanır ve Docker container'ına "
        "ortam değişkeni olarak iletilir; repoya commit edilmez. "
        "Kişisel veri içeren senaryolarda KVKK ve GDPR gereği şu önlemler "
        "alınmalıdır: (i) hangi alanların PII olduğunu belirten meta veri; "
        "(ii) log redaction politikası; (iii) Anthropic'e veri gönderimi "
        "için rıza akışı; (iv) saklama süresi tanımı ve otomatik silme.",
    )
    add_page_break(doc)


# ---------- BÖLÜM 6: TEST VE DEĞERLENDİRME ----------

def build_bolum6_test(doc):
    add_heading_1(doc, "BÖLÜM 6. TEST VE DEĞERLENDİRME")
    add_body(
        doc,
        "Sistem üç farklı düzeyde doğrulanmıştır: (1) birim testler, (2) "
        "uçtan uca entegrasyon testleri ve (3) baseline karşılaştırmalı "
        "değerlendirme. Tüm testler GitHub Actions üzerinde Python 3.10 ve "
        "3.11 matrisi altında otomatik koşmaktadır.",
    )
    add_subheading(doc, "6.1. Birim Testler")
    add_body(
        doc,
        "Toplam 67 birim test mock'larla yazılmıştır ve dış HTTP bağlantısı "
        "gerektirmez. Kapsadıkları konular: BaseAgent yaşam döngüsü, registry "
        "kayıt/keşif, TF-IDF embedding, trust scorer, deadlock detector, "
        "delegation manager, orchestrator decomposer ve delegator, no-code "
        "validator ve role builder, dört örnek ajanın araçları.",
    )
    add_subheading(doc, "6.2. Entegrasyon Testleri")
    add_body(
        doc,
        "12 entegrasyon testi gerçek uvicorn sunucularını daemon thread'lerde "
        "ayağa kaldırır, httpx ile gerçek HTTP istekleri gönderir. Test fixture'ı "
        "registry, dört ajan ve role builder'ı tek session içinde başlatır; "
        "her test ardından sunucular kapatılır. Bu yaklaşım birim mock'ların "
        "yakalayamayacağı serileştirme/CORS/zamanlama hatalarını ortaya çıkarır.",
    )
    add_subheading(doc, "6.3. Değerlendirme Protokolü")
    add_body(
        doc,
        "Akademik karşılaştırma için MNACP'nin yanı sıra iki baseline "
        "tanımlanmıştır:",
    )
    add_bullet(doc, "Statik Atama: Önceden tanımlı bir keyword→agent sözlüğü kullanır (örn. 'csv' → DataAgent). Dinamik keşif veya delegasyon yoktur.")
    add_bullet(doc, "Merkezi (Delegasyonsuz): Tüm araçlar tek bir merkezi ajan üzerindedir. Anahtar kelime eşleşmesi ile araç seçer; ajanlar arası delegasyon olmaz.")
    add_caption_over(doc, "Tablo 6.1. Değerlendirme senaryoları (scenarios.json).")
    add_table_simple(
        doc,
        ["ID", "Açıklama", "Beklenen Ajanlar", "Zorluk"],
        [
            ["S1", "CSV yükle, temizle, istatistik, web karşılaştır, rapor", "DataAgent, SearchAgent, AnalysisAgent", "Zor"],
            ["S2", "Web'den veri topla, özetle, istatistiksel özet çıkar", "SearchAgent, DataAgent", "Orta"],
            ["S3", "Sayı listesinde trend analizi ve değişim yüzdesi", "AnalysisAgent", "Kolay"],
            ["S4", "Hisse CSV → trend analizi → web haber → yatırım raporu", "DataAgent, AnalysisAgent, SearchAgent", "Zor"],
            ["S5", "İki veri setini karşılaştır ve raporla", "AnalysisAgent", "Kolay"],
        ],
    )
    add_body(
        doc,
        "Her sistem her senaryoyu 3 kez koşturmuştur (toplam 15 koşum/sistem). "
        "MNACP gerçek HTTP üzerinden orkestratöre çağrı yaparak Claude API ile "
        "çalışmıştır; baseline'lar kural tabanlı simülasyondur. Her koşumda "
        "şunlar ölçülür: tamamlama (başarı/başarısızlık), seçilen ajan listesi, "
        "delegasyon zinciri uzunluğu, uçtan uca gecikme.",
    )
    add_subheading(doc, "6.4. Sayısal Sonuçlar")
    add_caption_over(doc, "Tablo 6.2. Sistem başına özet metrikler (15 koşum/sistem).")
    add_table_simple(
        doc,
        ["Sistem", "Tamamlama", "Ajan Doğruluğu", "Ort. Delegasyon", "P50 Gecikme"],
        [
            ["MNACP (bizim)", "%100", "%96,7", "4,27 (derinlik 1)", "18.641 ms"],
            ["Merkezi (delegasyonsuz)", "%100", "%72,2", "3,6 araç", "Simülasyon (~0 ms)"],
            ["Statik Atama", "%60", "%30,0", "1 ajan", "Simülasyon (~0 ms)"],
        ],
    )
    add_body(
        doc,
        "Statik atama 5 senaryodan ikisinde (S3, S5) tamamen başarısız olmuştur "
        "— anahtar kelime eşleşmesi olmadığında DataAgent'e fallback yapması, "
        "AnalysisAgent gerektiren senaryolarda hatalı seçimle sonuçlanmıştır. "
        "Merkezi baseline tüm senaryoları tamamlamış olsa da %72,2 ajan seçim "
        "doğruluğu, MNACP'nin %96,7 doğruluğunun belirgin altındadır. MNACP'nin "
        "embedding tabanlı keşfi statik atamanın yaklaşık 3 katı, merkezi "
        "baseline'ın 1,34 katı doğruluğa ulaşmıştır.",
    )

    # Görseller
    add_caption_over(doc, "Şekil 6.1. Sistem karşılaştırması — görev tamamlama oranı.")
    add_image_centered(doc, EVAL_DIR / "completion_rate.png", width_cm=14)

    add_caption_over(doc, "Şekil 6.2. Senaryo bazlı performans ısı haritası.")
    add_image_centered(doc, EVAL_DIR / "scenario_heatmap.png", width_cm=14)

    add_caption_over(doc, "Şekil 6.3. Gecikme dağılımı (P50/P90/P99).")
    add_image_centered(doc, EVAL_DIR / "latency_distribution.png", width_cm=15)

    add_caption_over(doc, "Şekil 6.4. Ajan bazlı performans dağılımı (monitor sayfası).")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(
        "[Ekran görüntüsü için yer tutucu: /monitor sayfasındaki ajan "
        "bazlı yığılmış bar chart — DataAgent, AnalysisAgent, vb. için "
        "başarılı/hata dağılımı]"
    ).italic = True

    add_subheading(doc, "6.5. Bulguların Tartışılması")
    add_body(
        doc,
        "Sonuçlar üç temel bulguya işaret etmektedir. Birincisi, embedding "
        "tabanlı dinamik keşif anahtar kelime eşleşmesinin yapısal kısıtlarını "
        "(senaryo metninin tetikleyici kelimeyi içermemesi, terim eşanlamlısı "
        "kullanılması) aşarak doğru ajan seçimini sağlamıştır. İkincisi, peer "
        "delegasyon mekanizması orkestratör darboğazını ortadan kaldırarak "
        "ajan→ajan zincirlerinin doğal biçimde oluşmasına olanak tanımıştır. "
        "Üçüncüsü, MNACP'nin gecikmesi (P50 ≈ 18,6 saniye) gerçek Claude API "
        "çağrılarından kaynaklanmakta olup, simülasyon baseline'larıyla doğrudan "
        "karşılaştırılamaz; adil bir gecikme karşılaştırması için baseline'ların "
        "da gerçek API çağrıları üzerine kurulması gerekmektedir.",
    )
    add_subheading(doc, "6.6. Sınırlamalar")
    add_bullet(doc, "Senaryo havuzu 5 örnek ile sınırlıdır; daha geniş ve farklı zorluk dağılımına sahip bir havuz daha güvenilir genelleme sağlar.")
    add_bullet(doc, "Embedding boyutu (256) küçük korpus için yeterli olsa da çok daha geniş ajan kataloğunda yeniden değerlendirilmelidir.")
    add_bullet(doc, "Güven skoru başlangıçta tüm ajanlar için 1,0 alınır; soğuk başlangıç periyodunda ayrım gücü düşüktür.")
    add_bullet(doc, "Mevcut deadlock kontrolü tek-süreçlidir; coğrafi olarak dağıtık konuşlandırmada eşgüdüm için ek bir kilit servisi gerekir.")
    add_page_break(doc)


# ---------- BÖLÜM 7: SONUÇLAR VE ÖNERİLER ----------

def build_bolum7(doc):
    add_heading_1(doc, "BÖLÜM 7. SONUÇLAR VE ÖNERİLER")
    add_body(
        doc,
        "Bu çalışmada, MCP standardının açık bıraktığı çok-ajanlı koordinasyon "
        "boşluğunu doldurmaya yönelik bir protokol ve referans uygulama "
        "geliştirilmiştir. MNACP; çalışma zamanında embedding tabanlı semantik "
        "ajan keşfi, üstel ağırlıklı güven skoru, Kosaraju algoritması ile "
        "deadlock tespiti, LangGraph durum makineli orkestratör, no-code "
        "GenericAgent üretimi ve peer delegasyon zincirlerini bir arada "
        "sunmaktadır. Tüm bileşenler HTTP üzerinden gevşek bağlı çalışmakta, "
        "Docker Compose ile tek komutla konuşlandırılmakta ve GitHub Actions "
        "üzerinde sürekli test edilmektedir.",
    )
    add_body(
        doc,
        "5 senaryolu, 15 koşumluk değerlendirme protokolünde MNACP %100 "
        "tamamlama ve %96,7 ajan seçim doğruluğuyla statik atama (%30) ve "
        "merkezi baseline (%72,2) sistemlerini belirgin şekilde geride "
        "bırakmıştır. Bulgular, LLM tabanlı çoklu-ajan sistemlerinde dinamik "
        "keşfin ve matematiksel güvenlik mekanizmalarının pratik "
        "uygulanabilirliğini ortaya koymaktadır.",
    )
    add_subheading(doc, "7.1. İleri Çalışma Önerileri")
    add_body(doc, "Algoritma ve karar mekanizmaları:")
    add_bullet(doc, "Embedding modelinin TF-IDF + SVD yerine modern dense embeddings (örn. Voyage AI, OpenAI text-embedding-3-small) ile değiştirilmesi ve karşılaştırılması.")
    add_bullet(doc, "Güven skorunun ardışık hata cezasını ajan tipine göre uyarlayan adaptif öğrenme mekanizmaları.")
    add_bullet(doc, "Peer delegasyon kararlarının sezgisel kuralla değil, küçük bir sınıflandırıcı modelle verilmesi.")

    add_body(doc, "Ölçek ve dağıtık konuşlandırma:")
    add_bullet(doc, "Çok düğümlü konuşlandırma için dağıtık deadlock tespiti — Chandy–Misra–Haas tabanlı algoritmaların entegrasyonu.")
    add_bullet(doc, "Registry'nin tek nokta arıza riskini ortadan kaldıracak şekilde Raft veya benzeri bir konsensüs algoritmasıyla replikasyonu.")

    add_body(doc, "Güvenlik:")
    add_bullet(doc, "Ajanlar arası mTLS veya OAuth 2.0 client credentials akışı ile kimlik doğrulama.")
    add_bullet(doc, "CodeAgent için Linux namespace + seccomp-bpf tabanlı katı sandbox; AST kontrolünün de bypass edilmesini engelleyecek ek kontroller.")
    add_bullet(doc, "Prompt injection için kullanıcı metnini sistem promptundan ayıran XML delimiterları ve tool gating mekanizması.")
    add_bullet(doc, "Tedarik zinciri saldırılarına karşı SBOM üretimi, image imzalama (Sigstore/cosign) ve pip-audit/npm-audit'in CI'a entegrasyonu.")

    add_body(doc, "Değerlendirme:")
    add_bullet(doc, "Daha geniş ve çeşitli senaryo havuzu üzerinde uzun erimli değerlendirme; baseline'ların da gerçek API çağrılarıyla koşturulması.")
    add_bullet(doc, "Adversarial senaryolar — prompt injection, sahte kayıt, sandbox kaçışı denemelerinin sistematik olarak test edilmesi.")
    add_page_break(doc)


# ---------- KAYNAKLAR ----------

def build_kaynaklar(doc):
    add_heading_1(doc, "KAYNAKLAR")

    refs = [
        "Microsoft Research. AutoGen: Enabling Next-Gen LLM Applications via "
        "Multi-Agent Conversation. https://github.com/microsoft/autogen, "
        "Erişim Tarihi: 2026.",
        "CrewAI Inc. CrewAI: Cutting-edge Framework for Orchestrating "
        "Role-Playing, Autonomous AI Agents. https://github.com/crewAIInc/crewAI, "
        "Erişim Tarihi: 2026.",
        "LangChain AI. LangGraph: Building Stateful, Multi-Actor Applications "
        "with LLMs. https://github.com/langchain-ai/langgraph, Erişim Tarihi: 2026.",
        "Anthropic. Model Context Protocol Specification. "
        "https://modelcontextprotocol.io, Kasım 2024.",
        "Anthropic. Claude API Documentation. https://docs.anthropic.com, "
        "Erişim Tarihi: 2026.",
        "Deerwester, S., Dumais, S. T., Furnas, G. W., Landauer, T. K., Harshman, R. "
        "Indexing by Latent Semantic Analysis. Journal of the American Society for "
        "Information Science, 41(6):391-407, 1990.",
        "Pedregosa, F., Varoquaux, G., Gramfort, A., et al. Scikit-learn: Machine "
        "Learning in Python. Journal of Machine Learning Research, 12:2825-2830, 2011.",
        "Sharir, M. A Strong-Connectivity Algorithm and its Applications in Data "
        "Flow Analysis. Computers & Mathematics with Applications, 7(1):67-72, 1981.",
        "Chandy, K. M., Misra, J. The Drinking Philosophers Problem. ACM "
        "Transactions on Programming Languages and Systems, 6(4):632-646, 1984.",
        "Tiangolo. FastAPI Documentation. https://fastapi.tiangolo.com, "
        "Erişim Tarihi: 2026.",
        "Pydantic. Pydantic v2 Documentation. https://docs.pydantic.dev, "
        "Erişim Tarihi: 2026.",
        "Vercel Inc. Next.js Documentation. https://nextjs.org/docs, "
        "Erişim Tarihi: 2026.",
        "ReactFlow. React Flow: Build Node-Based UIs with React. "
        "https://reactflow.dev, Erişim Tarihi: 2026.",
        "Astral. Ruff: An Extremely Fast Python Linter and Code Formatter. "
        "https://github.com/astral-sh/ruff, Erişim Tarihi: 2026.",
        "Docker Inc. Docker Compose Documentation. https://docs.docker.com/compose, "
        "Erişim Tarihi: 2026.",
        "GitHub Inc. GitHub Actions Documentation. https://docs.github.com/actions, "
        "Erişim Tarihi: 2026.",
    ]
    for i, r in enumerate(refs, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-1.0)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(f"[{i}]\t")
        run.font.size = Pt(11)
        run.font.name = "Times New Roman"
        run = p.add_run(r)
        run.font.size = Pt(11)
        run.font.name = "Times New Roman"
    add_page_break(doc)


def build_ozgecmis(doc):
    add_heading_1(doc, "ÖZGEÇMİŞ")
    add_body(
        doc,
        f"{PLACEHOLDER_ADSOYAD}, [doğum tarihi yer tutucu] tarihinde "
        "[doğum yeri] doğdu. İlk, orta ve lise eğitimini [şehir] tamamladı. "
        "[Yıl] yılında Sakarya Üniversitesi Bilgisayar Mühendisliği Bölümü'nü "
        "kazandı. [Staj bilgileri yer tutucu]. SAÜ Bilgisayar Mühendisliği "
        "Bölümünden [yıl] yılında mezun olmuştur. İlgi alanları arasında "
        "yapay zeka, çok-ajanlı sistemler, dağıtık yazılım mimarileri ve "
        "açık kaynak yazılım geliştirme yer almaktadır.",
    )
    add_page_break(doc)


def build_tutanak(doc):
    add_heading_1(doc, "BSM 401 BİLGİSAYAR MÜHENDİSLİĞİ TASARIMI")
    add_centered_paragraph(doc, "DEĞERLENDİRME VE SÖZLÜ SINAV TUTANAĞI", bold=True, size=14, space_after=24)

    add_body(doc, "KONU: MNACP — MCP-Native Çok Ajanlı Koordinasyon Protokolü")
    add_body(doc, f"ÖĞRENCİ: {PLACEHOLDER_OGRENCI_NO} / {PLACEHOLDER_ADSOYAD}")
    doc.add_paragraph()

    rows = [
        ["Yazılı Çalışma", "", "", ""],
        ["Çalışma kılavuza uygun olarak hazırlanmış mı?", "x", "0-5", ""],
        ["Teknik Yönden", "", "", ""],
        ["Problemin tanımı yapılmış mı?", "x", "0-5", ""],
        ["Mimarinin blok şeması çizilerek açıklanmış mı?", "", "", ""],
        ["Birimler arası bilgi akışı modeli verilmiş mi?", "", "", ""],
        ["Yazılım gereksinim listesi oluşturulmuş mu?", "", "", ""],
        ["Kullanılan araçlar/teknolojiler anlatılmış mı?", "", "", ""],
        ["UML ile modelleme yapılmış mı?", "", "", ""],
        ["Veri tabanı kullanıldıysa kavramsal model verilmiş mi?", "", "", ""],
        ["İş-zaman çizelgesi ve maliyet analizi yapılmış mı?", "", "", ""],
        ["Sürüm denetim sistemi kullanılmış mı?", "", "", ""],
        ["Sistem testleri ve iyileştirme süreci belgelenmiş mi?", "", "", ""],
        ["Performans testi yapılmış mı?", "", "", ""],
        ["Yapılan işlerin zorluk derecesi", "x", "0-25", ""],
        ["Sözlü Sınav", "", "", ""],
        ["Yapılan sunum başarılı mı?", "x", "0-5", ""],
        ["Soruları yanıtlama yetkinliği", "x", "0-20", ""],
        ["Devam Durumu", "", "", ""],
        ["Dönem içi raporlar düzenli hazırlandı mı?", "x", "0-5", ""],
        ["Toplam", "", "", ""],
    ]
    add_table_simple(doc, ["Değerlendirme Konusu", "İstenen", "Not Aralığı", "Not"], rows)

    doc.add_paragraph()
    add_body(doc, f"DANIŞMAN: {PLACEHOLDER_DANISMAN}")
    add_body(doc, "DANIŞMAN İMZASI:")


# ---------- ANA AKIŞ ----------

def main():
    doc = Document()

    # Sayfa kenar boşlukları (Sakarya şablonu: sol 3.5, diğerleri 2.5)
    for section in doc.sections:
        set_margins(section)

    set_default_style(doc)

    build_cover(doc)
    build_onsoz(doc)
    build_icindekiler_placeholder(doc)
    build_kisaltmalar(doc)
    build_sekiller_listesi(doc)
    build_tablolar_listesi(doc)
    build_ozet(doc)

    build_bolum1(doc)
    build_bolum2(doc)
    build_bolum3(doc)
    build_bolum4_algoritmalar(doc)
    build_bolum5_guvenlik(doc)
    build_bolum6_test(doc)
    build_bolum7(doc)

    build_kaynaklar(doc)
    build_ozgecmis(doc)
    build_tutanak(doc)

    doc.save(OUTPUT)
    print(f"DOCX olusturuldu: {OUTPUT}")
    print(f"Boyut: {os.path.getsize(OUTPUT) / 1024:.1f} KB")


if __name__ == "__main__":
    main()
